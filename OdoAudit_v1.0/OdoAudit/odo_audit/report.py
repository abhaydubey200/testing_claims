"""Writes the Excel audit workbook."""
import os, random, io
from datetime import datetime
import pandas as pd
from PIL import Image as PImage, ImageDraw, ImageFont
try:
    import pillow_heif; pillow_heif.register_heif_opener()
except Exception:
    pass
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XImg
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L
from . import config as C

F = "Arial"
def f(**k): return Font(name=F, **k)
HF = PatternFill("solid", fgColor="1F3864"); OF = PatternFill("solid", fgColor="C55A11")
RED = PatternFill("solid", fgColor="FFC7CE"); GRN = PatternFill("solid", fgColor="C6EFCE"); AMB = PatternFill("solid", fgColor="FFEB9C")
INP = PatternFill("solid", fgColor="FCE4D6")
thin = Side(style="thin", color="BFBFBF"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
HR_OPTS = ["Correct", "Wrong - photo HAS timestamp + location", "Wrong - photo has NO stamp", "Unclear - need original photo"]
ACTIONS = {
    "Possible AI-edited photo": "Ask for the original photo from the phone; verify odometer physically",
    "Possible AI-generated photo": "Ask for the original photo from the phone; verify odometer physically",
    "Stamp looks edited": "Ask for the original photo; confirm the trip",
    "GPS outside India": "Ask for the original photo; confirm the trip",
    "Photo dated after the claim date": "Ask for proof of the trip on the claim date",
    "Photo dated after claim was submitted": "Ask for proof of the trip",
    "Same photo for START and END": "Reject or recheck the claim",
    "Reused photo": "Ask for a fresh START and END photo every day",
    "Stamp date does not match claim": "Ask the employee to explain",
    "END photo taken before START": "Ask the employee to explain",
    "Odometer reading goes backwards": "Check readings in the full-size photos",
    "Odometer END below START": "Check readings in the full-size photos; recheck distance claimed",
    "KM claimed above odometer": "Check readings in the full-size photos",
    "Edited with photo software": "Ask for the original photo",
    "Screenshot instead of photo": "Remind employee to upload the original camera photo",
    "Unusual file type": "Check the photo",
}


def _hdr(ws, cols, widths, h=34, input_cols=()):
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        x = ws.cell(1, c); x.font = f(bold=True, color="FFFFFF"); x.fill = OF if c in input_cols else HF; x.border = BD
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[1].height = h; ws.freeze_panes = "C2"


def _fmt(ws, r, n, wrap=()):
    for c in range(1, n + 1):
        x = ws.cell(r, c); x.font = f(); x.border = BD; x.alignment = Alignment(vertical="center", wrap_text=c in wrap)


def _thumb(path, size, tmpdir, name):
    try:
        im = PImage.open(path).convert("RGB"); im.thumbnail(size)
        out = os.path.join(tmpdir, name + ".jpg"); im.save(out, quality=78); return out
    except Exception:
        return None


def write_report(P: pd.DataFrame, claims: pd.DataFrame, odo: pd.DataFrame, meta: dict, out_path: str, tmpdir: str):
    os.makedirs(tmpdir, exist_ok=True)
    wb = Workbook()
    n = len(P); PL = n + 1; R = "'Photo Review'!"

    # ---------------- Photo Review ----------------
    ph = wb.active; ph.title = "Photo Review"
    cols = ["Photo #", "Employee_code", "User_Name", "Designation", "Zone", "HQ", "L1 Approver", "Claim_Id", "Claim_Date",
            "Claim Key", "Side", "Claim_Amount (₹)", "Photo", "App Result", "Stamp Date(s)", "Stamp Time", "Stamp GPS / App",
            "Odometer (OCR)", "Red Flags", "Flag Level", "Check Priority", "HR Check", "Final Result", "Compliant?",
            "Photo Link", "Lookup Key", "Text read from photo"]
    _hdr(ph, cols, [7, 11, 20, 8, 9, 12, 16, 9, 11, 8, 7, 10, 22, 22, 13, 8, 18, 11, 46, 8, 16, 24, 22, 9, 11, 10, 40], input_cols=(22,))
    for k, r in enumerate(P.itertuples()):
        row = k + 2
        gps = ", ".join(x for x in [r.latlon or "", r.stamp_app or ""] if x)
        odo_v = ", ".join(f"{v:g}" for v in (r.odo_candidates or [])[:3])
        lvl = {0: "High", 1: "Medium", 2: "Low"}.get(r.flag_priority, "")
        vals = [r.pid, r.emp, r.name, r.designation, r.zone, r.hq, r.l1, r.claim_id, r.claim_date, r.claim_key, r.side, r.amount, None,
                r.result, ", ".join(pd.to_datetime(d).strftime("%d-%b-%Y") for d in (r.stamp_dates or [])), r.stamp_time, gps, odo_v,
                r.flag_text, lvl, r.check_priority, None,
                f'=IF(V{row}="{HR_OPTS[1]}","{C.RESULT_OK}",IF(V{row}="{HR_OPTS[2]}","No stamp",N{row}))',
                f'=IF(W{row}="{C.RESULT_OK}","Yes","No")', "Open photo", f'=J{row}&"|"&K{row}', (r.ocr_text or "")[:500]]
        ph.append(vals); _fmt(ph, row, len(cols), wrap=(3, 7, 14, 15, 17, 19, 21, 22, 23, 27))
        ph.cell(row, 9).number_format = "dd-mmm-yyyy"; ph.cell(row, 12).number_format = "#,##0.00"
        lk = ph.cell(row, 25); lk.hyperlink = r.url; lk.font = f(color="0563C1", underline="single")
        if r.check_priority: ph.cell(row, 22).fill = INP
        ph.row_dimensions[row].height = 98
        if r.path:
            t = _thumb(r.path, C.THUMB_SIZE, tmpdir, f"p{r.pid}")
            if t: ph.add_image(XImg(t), f"M{row}")
    ph.auto_filter.ref = f"A1:{L(len(cols))}{PL}"
    dv = DataValidation(type="list", formula1='"' + ",".join(HR_OPTS) + '"', allow_blank=True); dv.add(f"V2:V{PL}"); ph.add_data_validation(dv)
    ph.conditional_formatting.add(f"W2:X{PL}", FormulaRule(formula=['$X2="No"'], fill=RED))
    ph.conditional_formatting.add(f"W2:X{PL}", FormulaRule(formula=['$X2="Yes"'], fill=GRN))
    ph.conditional_formatting.add(f"S2:T{PL}", FormulaRule(formula=['$T2="High"'], fill=RED))
    ph.conditional_formatting.add(f"S2:T{PL}", FormulaRule(formula=['$T2="Medium"'], fill=AMB))
    ph.conditional_formatting.add(f"V2:V{PL}", FormulaRule(formula=['LEFT(V2,5)="Wrong"'], fill=RED, font=Font(name=F, bold=True, color="9C0006")))
    ph.conditional_formatting.add(f"V2:V{PL}", FormulaRule(formula=['V2="Correct"'], fill=GRN))
    ph.column_dimensions["Z"].hidden = True

    # ---------------- Claim Review ----------------
    cl = wb.create_sheet("Claim Review", 0)
    ccols = ["Claim Key", "Employee_code", "User_Name", "Zone", "HQ", "L1 Approver", "Claim_Id", "Claim_Date", "Submit Date",
             "Claim_Amount (₹)", "START Result", "END Result", "Claim Result", "Odometer START (OCR)", "Odometer END (OCR)",
             "Odometer km (OCR)", "Photos with Red Flags"]
    _hdr(cl, ccols, [8, 11, 20, 9, 12, 16, 9, 11, 11, 10, 22, 22, 15, 10, 10, 9, 9])
    nc = len(claims); CL = nc + 1
    for k, c in enumerate(claims.itertuples()):
        row = k + 2
        vals = [c.claim_key, c.emp, c.name, c.zone, c.hq, c.l1, c.claim_id, c.claim_date, c.submit_date, c.amount,
                f'=IFERROR(INDEX({R}$W$2:$W${PL},MATCH($A{row}&"|START",{R}$Z$2:$Z${PL},0)),"Photo missing / not downloaded")',
                f'=IFERROR(INDEX({R}$W$2:$W${PL},MATCH($A{row}&"|END",{R}$Z$2:$Z${PL},0)),"Photo missing / not downloaded")',
                f'=IF(AND(K{row}="{C.RESULT_OK}",L{row}="{C.RESULT_OK}"),"Compliant","NON-COMPLIANT")',
                c.odo_start, c.odo_end, c.odo_km,
                f'=COUNTIFS({R}$J$2:$J${PL},A{row},{R}$S$2:$S${PL},"?*")']
        cl.append(vals); _fmt(cl, row, len(ccols), wrap=(3, 6, 11, 12))
        for cc in (8, 9): cl.cell(row, cc).number_format = "dd-mmm-yyyy"
        cl.cell(row, 10).number_format = "#,##0.00"
    cl.auto_filter.ref = f"A1:Q{CL}"
    cl.conditional_formatting.add(f"M2:M{CL}", FormulaRule(formula=['$M2="NON-COMPLIANT"'], fill=RED, font=Font(name=F, bold=True, color="9C0006")))
    cl.conditional_formatting.add(f"M2:M{CL}", FormulaRule(formula=['$M2="Compliant"'], fill=GRN))
    cl.conditional_formatting.add(f"K2:L{CL}", FormulaRule(formula=[f'AND(K2<>"",K2<>"{C.RESULT_OK}")'], fill=RED))
    cl.conditional_formatting.add(f"Q2:Q{CL}", FormulaRule(formula=["Q2>0"], fill=AMB))
    CR = "'Claim Review'!"

    # ---------------- Employee Summary ----------------
    es = wb.create_sheet("Employee Summary", 0)
    emp = (P.groupby("emp").agg(name=("name", "first"), des=("designation", "first"), zone=("zone", "first"), hq=("hq", "first"),
                                l1=("l1", "first"), ok=("result", lambda s: (s == C.RESULT_OK).sum()), n=("result", "size"),
                                fl=("flag_priority", "min")).reset_index())
    emp["rank"] = emp.apply(lambda r: 0 if r.ok == 0 else (1 if r.ok < r.n else 2), axis=1)
    emp = emp.sort_values(["rank", "fl", "emp"]).reset_index(drop=True)
    efl = {}
    for e, g in P.groupby("emp"):
        cats = {}
        for fl in g["flag_list"]:
            for p, cat, _ in fl: cats[cat] = cats.get(cat, 0) + 1
        efl[e] = " | ".join(f"{k} ({v})" for k, v in sorted(cats.items(), key=lambda x: -x[1]))
    ecols = ["Employee_code", "User_Name", "Designation", "Zone", "HQ", "L1 Approver", "Odometer Claims", "Photos",
             "Photos with Timestamp + Location", "Photos NOT Compliant", "Non-Compliant Claims", "Amount on Non-Compliant Claims (₹)",
             "Total Odometer Claim Amount (₹)", "High Red Flags", "Employee Status", "Red Flags (photos)"]
    _hdr(es, ecols, [11, 20, 9, 9, 12, 16, 9, 8, 11, 10, 10, 13, 13, 8, 22, 55], h=48)
    ne = len(emp); EL = ne + 1
    for k, r in emp.iterrows():
        row = k + 2
        vals = [r.emp, r["name"], r.des, r.zone, r.hq, r.l1,
                f"=COUNTIF({CR}$B$2:$B${CL},A{row})", f"=COUNTIF({R}$B$2:$B${PL},A{row})",
                f'=COUNTIFS({R}$B$2:$B${PL},A{row},{R}$X$2:$X${PL},"Yes")', f"=H{row}-I{row}",
                f'=COUNTIFS({CR}$B$2:$B${CL},A{row},{CR}$M$2:$M${CL},"NON-COMPLIANT")',
                f'=SUMIFS({CR}$J$2:$J${CL},{CR}$B$2:$B${CL},A{row},{CR}$M$2:$M${CL},"NON-COMPLIANT")',
                f"=SUMIFS({CR}$J$2:$J${CL},{CR}$B$2:$B${CL},A{row})",
                f'=COUNTIFS({R}$B$2:$B${PL},A{row},{R}$T$2:$T${PL},"High")',
                f'=IF(I{row}=0,"No stamp on any photo",IF(J{row}>0,"Some photos without stamp","Compliant"))', efl.get(r.emp, "")]
        es.append(vals); _fmt(es, row, len(ecols), wrap=(2, 6, 15, 16))
        for cc in (12, 13): es.cell(row, cc).number_format = "#,##0.00"
    es.auto_filter.ref = f"A1:P{EL}"
    es.conditional_formatting.add(f"O2:O{EL}", FormulaRule(formula=['$O2="No stamp on any photo"'], fill=RED, font=Font(name=F, bold=True, color="9C0006")))
    es.conditional_formatting.add(f"O2:O{EL}", FormulaRule(formula=['$O2="Some photos without stamp"'], fill=AMB))
    es.conditional_formatting.add(f"O2:O{EL}", FormulaRule(formula=['$O2="Compliant"'], fill=GRN))
    es.conditional_formatting.add(f"N2:N{EL}", FormulaRule(formula=["N2>0"], fill=RED))
    ES = "'Employee Summary'!"

    # ---------------- Red Flags ----------------
    rf = wb.create_sheet("Red Flags")
    rows = []
    for (e, cat), g in pd.DataFrame([(r.emp, cat, pr, r.pid, r.claim_id, t) for r in P.itertuples() for pr, cat, t in r.flag_list],
                                    columns=["emp", "cat", "pr", "pid", "claim", "txt"]).groupby(["emp", "cat"]) if any(len(x) for x in P["flag_list"]) else []:
        pr = min(g.pr, key=lambda p: ["High", "Medium", "Low"].index(p))
        det = " | ".join(dict.fromkeys(f"#{a}: {t}" for a, t in zip(g.pid, g.txt)))
        rows.append((pr, cat, e, P[P.emp == e].name.iloc[0], P[P.emp == e].l1.iloc[0], ", ".join(f"#{x}" for x in dict.fromkeys(g.pid)),
                     ", ".join(str(x) for x in dict.fromkeys(g.claim)), det[:1500], ACTIONS.get(cat, "Check the photos")))
    rows.sort(key=lambda r: (["High", "Medium", "Low"].index(r[0]), r[1], r[2]))
    rcols = ["Priority", "Category", "Employee_code", "User_Name", "L1 Approver", "Photo #", "Claim_Id(s)", "Details", "Suggested Action"]
    _hdr(rf, rcols, [9, 24, 11, 20, 16, 14, 16, 70, 30], h=28); rf.freeze_panes = "D2"
    for k, rr in enumerate(rows):
        rf.append(list(rr)); _fmt(rf, k + 2, len(rcols), wrap=(2, 4, 5, 6, 7, 8, 9))
    RL = max(2, len(rows) + 1)
    rf.auto_filter.ref = f"A1:I{RL}"
    rf.conditional_formatting.add(f"A2:A{RL}", FormulaRule(formula=['$A2="High"'], fill=RED, font=Font(name=F, bold=True, color="9C0006")))
    rf.conditional_formatting.add(f"A2:A{RL}", FormulaRule(formula=['$A2="Medium"'], fill=AMB))

    # ---------------- Corner Check (AI mark visual scan) ----------------
    cc = wb.create_sheet("AI Mark Corner Check")
    cc["A1"] = "Bottom-right corner of every photo. Look for a small four-point sparkle (✦) - the visible mark Google Gemini adds to AI-made or AI-edited images. Note any photo # you see and check it in Photo Review."
    cc["A1"].font = f(bold=True); cc["A1"].alignment = Alignment(wrap_text=True); cc.merge_cells("A1:L1"); cc.row_dimensions[1].height = 34
    font = None
    for fp in ("arialbd.ttf", "C:/Windows/Fonts/arialbd.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"):
        try: font = ImageFont.truetype(fp, 14); break
        except Exception: pass
    S, per = 120, 60; okp = P[P.path.notna()]
    rowpos = 3
    for s in range(0, len(okp), per):
        chunk = okp.iloc[s:s + per]; sheet = PImage.new("RGB", (10 * S, ((len(chunk) + 9) // 10) * (S + 16)), "white"); d = ImageDraw.Draw(sheet)
        for k, r in enumerate(chunk.itertuples()):
            try:
                im = PImage.open(r.path).convert("RGB"); w, h = im.size
                c = im.crop((int(w * .78), int(h * .82), w, h)).resize((S, S))
                x, y = (k % 10) * S, (k // 10) * (S + 16); sheet.paste(c, (x, y + 16)); d.text((x + 2, y), f"#{r.pid}", fill="red", font=font)
            except Exception: pass
        pth = os.path.join(tmpdir, f"corner{s}.png"); sheet.save(pth)
        xi = XImg(pth); cc.add_image(xi, f"A{rowpos}"); rowpos += int(sheet.height / 20) + 2

    # ---------------- Raw Data with Photos ----------------
    rd = wb.create_sheet("Raw Data with Photos")
    raw_cols = [c for c in odo.columns if not c.startswith("_") and not c.endswith("_num")]
    rd.append(raw_cols)
    for c in range(1, len(raw_cols) + 1):
        x = rd.cell(1, c); x.font = f(bold=True, color="FFFFFF"); x.fill = HF; x.alignment = Alignment(wrap_text=True, vertical="center")
        rd.column_dimensions[L(c)].width = 14
    si, ei = raw_cols.index(C.COL_START) + 1, raw_cols.index(C.COL_END) + 1
    for col in (si, ei): rd.column_dimensions[L(col)].width = 22
    pk = {(r.claim_key, r.side): r for r in P.itertuples()}
    for k, r in enumerate(odo.itertuples(index=False)):
        row = k + 2; rec = r._asdict()
        rd.append([odo.iloc[k][c] for c in raw_cols])
        rd.row_dimensions[row].height = 98
        for side, col in (("START", si), ("END", ei)):
            p = pk.get((int(odo.iloc[k]["_row"]), side))
            cell = rd.cell(row, col)
            if p is not None and p.path:
                t = os.path.join(tmpdir, f"p{p.pid}.jpg")
                if os.path.exists(t):
                    cell.value = None; rd.add_image(XImg(t), f"{L(col)}{row}"); cell.hyperlink = p.url
            elif cell.value:
                cell.fill = RED
    rd.freeze_panes = "E2"; rd.auto_filter.ref = f"A1:{L(len(raw_cols))}{len(odo) + 1}"

    # ---------------- Summary ----------------
    sm = wb.create_sheet("Summary", 0)
    sm.column_dimensions["A"].width = 60; sm.column_dimensions["B"].width = 18
    def put(t, v=None, bold=False, size=11, fmt=None):
        sm.append([t, v]); rr = sm.max_row
        sm.cell(rr, 1).font = f(bold=bold, size=size); sm.cell(rr, 2).font = f(bold=True)
        if v is not None: sm.cell(rr, 1).border = BD; sm.cell(rr, 2).border = BD
        if fmt: sm.cell(rr, 2).number_format = fmt
        return rr
    put("Odometer Photo Audit", bold=True, size=14)
    put(f"Source file: {meta['input']}   |   Run: {meta['run_time']}   |   App version {meta['version']}")
    put("A photo is compliant only if a date/time AND a location (address or lat/long) are printed on it. Results were read automatically "
        "from the photos (offline); check marked photos in Photo Review column V before sending.")
    for r_ in (2, 3): sm.merge_cells(start_row=r_, start_column=1, end_row=r_, end_column=5); sm.cell(r_, 1).alignment = Alignment(wrap_text=True); sm.row_dimensions[r_].height = 30
    put("")
    put("Photos", bold=True)
    put("Odometer photos in file", f"=COUNTA({R}A2:A{PL})")
    for k in C.RESULTS: put(f"   {k}", f'=COUNTIF({R}W2:W{PL},"{k}")')
    put("Photos NOT compliant", f'=COUNTIF({R}X2:X{PL},"No")', bold=True)
    put("Photos with a red flag", f'=COUNTIF({R}S2:S{PL},"?*")')
    put("")
    put("Claims", bold=True)
    put("Odometer claims", f"=COUNTA({CR}A2:A{CL})")
    put("Compliant claims", f'=COUNTIF({CR}M2:M{CL},"Compliant")')
    put("NON-COMPLIANT claims", f'=COUNTIF({CR}M2:M{CL},"NON-COMPLIANT")', bold=True)
    put("Amount on non-compliant claims (₹)", f'=SUMIF({CR}M2:M{CL},"NON-COMPLIANT",{CR}J2:J{CL})', fmt="#,##0.00")
    put("Total odometer claim amount (₹)", f"=SUM({CR}J2:J{CL})", fmt="#,##0.00")
    put("")
    put("Employees", bold=True)
    put("Employees", f"=COUNTA({ES}A2:A{EL})")
    a = put("   No stamp on any photo", f'=COUNTIF({ES}O2:O{EL},"No stamp on any photo")')
    b = put("   Some photos without stamp", f'=COUNTIF({ES}O2:O{EL},"Some photos without stamp")')
    put("   Compliant", f'=COUNTIF({ES}O2:O{EL},"Compliant")')
    put("Employees with at least one non-compliant photo", f"=B{a}+B{b}", bold=True)
    put("Employees with a High red flag", f'=COUNTIF({ES}N2:N{EL},">0")')
    put("")
    put("Red flags (see Red Flags sheet)", bold=True)
    for k in ("High", "Medium", "Low"): put(f"   {k} priority findings", f"=COUNTIF('Red Flags'!A2:A{RL},\"{k}\")")
    put("")
    put("HR verification before sending", bold=True)
    t = put("Photos marked for HR check (Photo Review column U)", f'=COUNTIF({R}U2:U{PL},"?*")')
    c_ = put("   Checked so far", f'=COUNTIFS({R}U2:U{PL},"?*",{R}V2:V{PL},"?*")')
    put("   Still to check", f"=B{t}-B{c_}")
    put("Results found WRONG by HR", f'=COUNTIF({R}V2:V{PL},"Wrong*")', bold=True)
    put("Random-sample results found wrong", f'=COUNTIFS({R}U2:U{PL},"3*",{R}V2:V{PL},"Wrong*")')
    put("")
    put("Run details", bold=True)
    for k, v in meta["log"]: put(k, v)
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)


def check_priorities(P: pd.DataFrame, seed=2026):
    """Which photos HR should look at before sending: red flags, likely misses, and a random sample."""
    pr = {}
    for r in P.itertuples():
        if r.flag_priority <= 1: pr[r.pid] = "1 - Red flag"
    ok_by_emp = P.groupby("emp").result.apply(lambda s: (s == C.RESULT_OK).sum())
    for r in P.itertuples():
        if r.pid not in pr and r.result != C.RESULT_OK and ok_by_emp.get(r.emp, 0) > 0:
            pr[r.pid] = "2 - Employee uses stamps elsewhere"
    rnd = random.Random(seed)
    nos = [r.pid for r in P.itertuples() if r.pid not in pr and r.result != C.RESULT_OK]
    yes = [r.pid for r in P.itertuples() if r.pid not in pr and r.result == C.RESULT_OK]
    for p in rnd.sample(nos, min(len(nos), max(10, len(nos) // 8))): pr[p] = "3 - Random sample (no stamp)"
    for p in rnd.sample(yes, min(len(yes), max(5, len(yes) // 10))): pr[p] = "3 - Random sample (stamped)"
    return P.pid.map(lambda p: pr.get(p, ""))
